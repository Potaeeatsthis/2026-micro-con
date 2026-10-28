#!/usr/bin/env python3
"""Offline regression check for Robotic_Arm_Kinematics_Workbook.xlsx.

Owner: IK / mechanical member.

Checks the workbook's candidate signs, per-joint endpoint table, FK/IK input
validation, the q4 = 0 IK endpoint summary, the preserved typed CAD vectors, and
the physical repeat/return statistics on Object_Check. When LibreOffice is
available, formulas are recalculated in a temporary copy, so the test never
trusts a stale cached value and never edits the workbook. No hardware is
commanded.

Run: python3 ProjectDevGuide/tests/host/test_kinematics_workbook.py
"""
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parents[2] / "docs" / "Robotic_Arm_Kinematics_Workbook.xlsx"
LO = shutil.which("libreoffice") or shutil.which("soffice")
NEEDS_LO = unittest.skipUnless(LO, "LibreOffice is not installed")


def recalculate(mutate=None):
    """Return a data_only workbook recalculated by LibreOffice in a temp dir."""
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        source = root / "source"
        out = root / "out"
        profile = root / "profile"
        source.mkdir()
        out.mkdir()
        profile.mkdir()

        wb = openpyxl.load_workbook(WORKBOOK)
        if mutate:
            mutate(wb)
        wb.calculation.fullCalcOnLoad = True
        path = source / "workbook.xlsx"
        wb.save(path)

        env = os.environ.copy()
        env["HOME"] = str(profile)
        subprocess.run(
            [
                LO,
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out),
                str(path),
            ],
            env=env,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        return openpyxl.load_workbook(out / "workbook.xlsx", data_only=True)


class WorkbookStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw = openpyxl.load_workbook(WORKBOOK)

    def test_signs_zeros_and_typed_cad_data_are_preserved(self):
        inputs = self.raw["Inputs"]
        self.assertEqual([inputs[f"B{r}"].value for r in range(27, 32)], [1, -1, 1, 1, -1])
        self.assertEqual([inputs[f"B{r}"].value for r in range(12, 17)], [85, 145, 88, 88, 96])
        self.assertEqual(inputs["B19"].value, 20)
        self.assertEqual([inputs[f"B{r}"].value for r in range(22, 25)], [10.15, 9.52, 438.5])

        objects = self.raw["Object_Check"]
        self.assertEqual([objects[c].value for c in ("B5", "C5", "D5", "B8", "C8", "D9")],
                         [10.15, 9.52, 438.5, 10.1, 9.37, 438.6])

        vectors = self.raw["Test_Vectors"]
        self.assertEqual(
            [[vectors.cell(r, c).value for c in range(7, 10)] for r in range(5, 14)],
            [
                [10.15, 9.52, 438.5],
                [13.55, 3.17, 438.5],
                [-160.25, 9.52, 385.88],
                [127.47, 9.53, 414.03],
                [9.03, 16.78, 438.5],
                [-40.49, 9.52, 417.97],
                [34.17, 48.76, 422.31],
                [266.49, -214.96, 118.63],
                [-94.04, 45.03, 338.71],
            ],
        )
        self.assertEqual(
            [[vectors.cell(r, c).value for c in range(11, 14)] for r in range(19, 23)],
            [
                [179.92, 30.08, 99.98],
                [179.99, 30.09, 99.9],
                [150.03, -40.15, 199.98],
                [99.77, -0.04, 300.19],
            ],
        )

        model = self.raw["Model_Params"]
        self.assertEqual([model[c].value for c in ("B5", "C5", "D5")],
                         ["=Joint_Axes!L6", "=Joint_Axes!M8", "=Joint_Axes!N6"])
        axes = self.raw["Joint_Axes"]
        self.assertEqual([axes[c].value for c in ("F6", "G6", "H6", "F9", "G9", "H9")],
                         [-23.15, 59.72, 14.08, -23.14, 299.24, 13.18])

    def test_per_joint_endpoints_and_checks(self):
        inputs = self.raw["Inputs"]
        self.assertEqual(
            [[inputs[f"A{r}"].value, inputs[f"B{r}"].value, inputs[f"C{r}"].value] for r in range(35, 41)],
            [
                ["J1", 0, 180],
                ["J2", 50, 180],
                ["J3", 45, 180],
                ["J4", 0, 180],
                ["J5", 25, 140],
                ["J6", 0, 60],
            ],
        )
        fk = self.raw["FK_Calculator"]
        for row in range(5, 10):
            self.assertIn(f"Inputs!$B${row + 30}", fk[f"E{row}"].value)
            self.assertIn(f"Inputs!$C${row + 30}", fk[f"E{row}"].value)
        ik = self.raw["IK_Calculator"]
        for row in range(14, 19):
            self.assertIn(f"Inputs!$B${row + 21}", ik[f"D{row}"].value)
            self.assertIn(f"Inputs!$C${row + 21}", ik[f"D{row}"].value)
        self.assertEqual(fk["D10"].value, "Endpoint summary")
        for token in ("Inputs!$B$32", "COUNTIF(E5:E9", "B6<90", "B9<65"):
            self.assertIn(token, fk["E10"].value)

    def test_sign_and_branch_validation_is_explicit(self):
        self.assertIn("OK: signs are ±1", self.raw["Inputs"]["B32"].value)
        self.assertIn("exactly +1", self.raw["IK_Calculator"]["B11"].value)
        self.assertIn("servo signs", self.raw["IK_Calculator"]["B11"].value)
        for sheet, target in (("Inputs", "B27:B31"), ("IK_Calculator", "B9")):
            validations = self.raw[sheet].data_validations.dataValidation
            found = [dv for dv in validations if target in str(dv.sqref)]
            self.assertEqual(len(found), 1, (sheet, target))
            self.assertTrue(found[0].showErrorMessage)
            self.assertEqual(found[0].errorStyle, "stop")

    def test_physical_repeats_and_rotation_formulas_are_recorded(self):
        objects = self.raw["Object_Check"]
        self.assertEqual([objects[c].value for c in ("B19", "B20", "B21", "C19", "C20", "C21")],
                         [6, 7.5, 7, 7, 10, 9])
        self.assertEqual(objects["B22"].value, "=AVERAGE(B19:B21)")
        self.assertEqual(objects["C22"].value, "=AVERAGE(C19:C21)")
        self.assertEqual(objects["B23"].value, "=MIN(B19:B21)")
        self.assertEqual(objects["C23"].value, "=MIN(C19:C21)")
        self.assertEqual(objects["B24"].value, "=MAX(B19:B21)")
        self.assertEqual(objects["C24"].value, "=MAX(C19:C21)")
        self.assertEqual(objects["B25"].value, "=B24-B23")
        self.assertEqual(objects["C25"].value, "=C24-C23")
        self.assertEqual([objects[c].value for c in ("B29", "C29", "D29", "B33", "C33")],
                         [6, 8, 10, 8, 7])
        self.assertEqual(objects["B34"].value, "=B33-B22")
        self.assertEqual(objects["C34"].value, "=C33-C22")
        self.assertEqual(objects["D35"].value, 446)
        self.assertEqual(objects["D36"].value, "=D35-15/2")
        self.assertEqual(objects["B42"].value, "=Inputs!B22")
        self.assertEqual(objects["B44"].value, "=B22")
        self.assertIn("RADIANS(15)", objects["B43"].value)
        self.assertIn("RADIANS(15)", objects["B45"].value)
        self.assertIn("not independent Z verification", objects["A48"].value)
        self.assertIn("static measurement repeatability", objects["A15"].value)
        self.assertIn("not grip repeatability", objects["A15"].value)
        self.assertIn("outside the starting X range", objects["F34"].value)

    def test_fk_pitch_is_planar_only(self):
        fk = self.raw["FK_Calculator"]
        self.assertIn("only when q4 = 0", fk["A21"].value)
        self.assertIn("Planar only: q4 must be 0", fk["B21"].value)

    def test_notes_do_not_overclaim(self):
        readme = self.raw["README"]
        self.assertIn("does not command hardware", readme["A2"].value)
        self.assertIn("hardware measurement", readme["A19"].value)

        servo_map = self.raw["Servo_Map"]
        self.assertIn("provisional", servo_map["A20"].value)
        self.assertIn("by construction", servo_map["A22"].value)
        self.assertIn("do not prove", servo_map["A23"].value)

        vectors = self.raw["Test_Vectors"]
        self.assertIn("0.028 mm", vectors["A24"].value)
        self.assertIn("not hardware proof", vectors["A24"].value)
        self.assertIn("9 poses", self.raw["Model_Params"]["A23"].value)
        self.assertIn("0.01512 mm", self.raw["Model_Params"]["A23"].value)

        objects = self.raw["Object_Check"]
        self.assertEqual([objects[c].value for c in ("B19", "B20", "B21", "C19", "C20", "C21")],
                         [6, 7.5, 7, 7, 10, 9])
        self.assertEqual([objects[c].value for c in ("B29", "C29", "D29", "B33", "C33", "D35")],
                         [6, 8, 10, 8, 7, 446])
        self.assertIn("Level-top assumption", objects["F36"].value)
        self.assertIn("not an independently measured", objects["F5"].value)

        inputs = self.raw["Inputs"]
        self.assertIn("20", inputs["D19"].value)
        self.assertNotIn("Unknown", inputs["D19"].value)
        self.assertNotIn("blank", inputs["D19"].value)
        self.assertNotIn("unknown/blank", self.raw["Servo_Map"]["A12"].value.lower())
        self.assertIn("20", self.raw["Open_Items"]["C9"].value)
        self.assertNotIn("blank", self.raw["Open_Items"]["D9"].value)
        self.assertIn("grip/TCP", self.raw["Open_Items"]["D12"].value)
        self.assertIn("not collision-certified", self.raw["Open_Items"]["C11"].value)
        for cell in (readme["A24"], self.raw["Inputs"]["A41"],
                     self.raw["IK_Calculator"]["D12"], self.raw["Open_Items"]["D11"]):
            self.assertIn("commanded-state and pending-endpoint interlocks", cell.value)
            self.assertIn("no automated collision guarantee", cell.value)
            self.assertNotIn("must still check the current and pending paths", cell.value)


@NEEDS_LO
class WorkbookFormulaResults(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw = openpyxl.load_workbook(WORKBOOK)
        cls.calc = recalculate()

    def test_default_ik_flags_j2_without_hiding_the_target(self):
        ik = self.calc["IK_Calculator"]
        self.assertEqual(ik["B11"].value, "REACHABLE (geometry only)")
        self.assertIn("OUTSIDE installed servo endpoints", ik["B12"].value)
        self.assertEqual(ik["D15"].value, "OUT OF RANGE")
        self.assertAlmostEqual(ik["C15"].value, 185.724554421762, places=9)

        self.assertEqual([ik[c].value for c in ("B5", "B6", "B7", "B8")], [120, 20, 300, 90])
        self.assertEqual(ik["B17"].value, 0)
        self.assertAlmostEqual(ik["C17"].value, 88.0, places=9)
        self.assertAlmostEqual(ik["C18"].value, 38.1070130457768, places=9)
        self.assertEqual(self.calc["FK_Calculator"]["E10"].value, "RANGES OK; J2/J5 coupling OK")

    def test_physical_repeat_statistics_and_15deg_rotation(self):
        obj = self.calc["Object_Check"]
        self.assertAlmostEqual(obj["B22"].value, 20.5 / 3, places=9)
        self.assertAlmostEqual(obj["C22"].value, 26 / 3, places=9)
        self.assertEqual([obj[c].value for c in ("B23", "C23", "B24", "C24")], [6, 7, 7.5, 10])
        self.assertAlmostEqual(obj["B25"].value, 1.5, places=9)
        self.assertAlmostEqual(obj["C25"].value, 3.0, places=9)
        self.assertEqual([obj[c].value for c in ("B29", "C29", "D29", "B33", "C33")], [6, 8, 10, 8, 7])
        self.assertAlmostEqual(obj["B34"].value, 7 / 6, places=9)
        self.assertAlmostEqual(obj["C34"].value, -5 / 3, places=9)
        self.assertAlmostEqual(obj["B33"].value - obj["B24"].value, 0.5, places=9)
        self.assertAlmostEqual(obj["D36"].value, 438.5, places=9)
        self.assertAlmostEqual(obj["D37"].value, 438.5, places=9)
        self.assertAlmostEqual(obj["D38"].value, 0.0, places=9)
        self.assertAlmostEqual(obj["B43"].value, 7.340189827458046, places=9)
        self.assertAlmostEqual(obj["C43"].value, 11.822627174062516, places=9)
        self.assertAlmostEqual(obj["B45"].value, 4.3573947554201204, places=9)
        self.assertAlmostEqual(obj["C45"].value, 10.139953969372483, places=9)
        self.assertAlmostEqual(obj["B46"].value, 2.9827950720379253, places=9)
        self.assertAlmostEqual(obj["C46"].value, 1.6826732046900332, places=9)

    def test_invalid_branch_and_sign_are_rejected(self):
        def bad_branch(wb):
            wb["IK_Calculator"]["B9"] = 0

        result = recalculate(bad_branch)
        self.assertTrue(result["IK_Calculator"]["B11"].value.startswith("INVALID INPUT: elbow branch"))
        self.assertEqual(result["IK_Calculator"]["B12"].value, "not evaluated: IK is not reachable")

        def bad_sign(wb):
            wb["Inputs"]["B28"] = 2

        result = recalculate(bad_sign)
        self.assertTrue(result["Inputs"]["B32"].value.startswith("INVALID"))
        self.assertTrue(result["IK_Calculator"]["B11"].value.startswith("INVALID INPUT: servo signs"))
        self.assertTrue(result["FK_Calculator"]["B10"].value.startswith("INVALID"))
        self.assertTrue(result["FK_Calculator"]["E10"].value.startswith("INVALID SIGNS"))

    def test_out_of_reach_is_not_silently_clamped(self):
        self.assertNotIn("MAX(-1,MIN(1", self.raw["IK_Calculator"]["B31"].value)
        self.assertIn("ACOS(B30)", self.raw["IK_Calculator"]["B31"].value)

        def too_far(wb):
            wb["IK_Calculator"]["B5"] = 1000

        result = recalculate(too_far)
        ik = result["IK_Calculator"]
        self.assertEqual(ik["B11"].value, "OUT OF REACH: arm cannot span target")
        self.assertIsNone(ik["B31"].value)
        self.assertIsNone(ik["B14"].value)

    def test_coupling_endpoint_is_not_collision_certification(self):
        def target(wb):
            ik = wb["IK_Calculator"]
            ik["B5"], ik["B6"], ik["B7"], ik["B8"], ik["B9"] = 155, -140, 365, 45, -1
            wb["FK_Calculator"]["B6"], wb["FK_Calculator"]["B9"] = 89.5, 64.0

        result = recalculate(target)
        ik = result["IK_Calculator"]
        self.assertIn("Within installed servo endpoints", ik["B12"].value)
        self.assertIn("J2/J5 coupling VIOLATED", ik["B12"].value)
        self.assertEqual(ik["D15"].value, "OK")
        self.assertEqual(ik["D18"].value, "OK")
        self.assertIn("J2/J5 coupling VIOLATED", result["FK_Calculator"]["E10"].value)

        def boundary(q2, q5):
            def mutate(wb):
                wb["IK_Calculator"]["C15"], wb["IK_Calculator"]["C18"] = q2, q5
                wb["FK_Calculator"]["B6"], wb["FK_Calculator"]["B9"] = q2, q5

            result = recalculate(mutate)
            return result["IK_Calculator"]["B12"].value, result["FK_Calculator"]["E10"].value

        for q2, q5, expected in ((90, 64.9, "coupling OK"), (89.9, 65, "coupling OK"),
                                 (89.9, 64.9, "coupling VIOLATED")):
            ik_summary, fk_summary = boundary(q2, q5)
            self.assertIn(expected, ik_summary)
            self.assertIn(expected, fk_summary)


if __name__ == "__main__":
    unittest.main()
